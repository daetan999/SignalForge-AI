"""Deterministic parsers for locally uploaded customer documents."""
from __future__ import annotations

import csv
import io
from pathlib import Path
from typing import Iterable

from docx import Document

from schemas.opportunity import ParsedDocument


def parse_document(name: str, payload: bytes) -> ParsedDocument:
    """Parse TXT, Markdown, CSV, or DOCX bytes into plain text."""
    suffix = Path(name).suffix.lower()
    if suffix in {".txt", ".md"}:
        text = payload.decode("utf-8", errors="replace")
        media_type = "text/plain"
    elif suffix == ".csv":
        decoded = payload.decode("utf-8-sig", errors="replace")
        rows = csv.reader(io.StringIO(decoded))
        text = "\n".join(" | ".join(cell.strip() for cell in row) for row in rows)
        media_type = "text/csv"
    elif suffix == ".docx":
        document = Document(io.BytesIO(payload))
        text = "\n".join(
            paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()
        )
        media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    else:
        raise ValueError(f"Unsupported file type: {suffix or 'no extension'}")

    if not text.strip():
        raise ValueError(f"{name} did not contain readable text.")
    return ParsedDocument(name=name, media_type=media_type, text=text.strip())


def parse_documents(files: Iterable[tuple[str, bytes]]) -> list[ParsedDocument]:
    """Parse a collection of named byte payloads."""
    parsed = [parse_document(name, payload) for name, payload in files]
    if not parsed:
        raise ValueError("At least one customer document is required.")
    return parsed
