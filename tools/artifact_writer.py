"""Create portable handoff artifacts from a validated analysis."""

from __future__ import annotations

import io
import json

from docx import Document

from schemas.opportunity import AnalysisResult


def analysis_to_markdown(result: AnalysisResult) -> str:
    lines = [
        f"# {result.profile.customer_name} — AI Opportunity Brief",
        "",
        f"**Use case:** {result.profile.use_case}",
        f"**Discovery coverage:** {result.coverage.score}%",
        f"**Readiness score:** {result.readiness_score}%",
        f"**Recommended action:** {result.decision.recommendation}",
        "",
        "## Confirmed Requirements",
    ]
    lines.extend(
        f"- {item.requirement}"
        for item in result.profile.requirements
        if item.status == "confirmed"
    )
    lines += ["", "## Discovery Gaps"]
    lines.extend(f"- [{gap.priority.upper()}] {gap.question}" for gap in result.gaps)
    lines += ["", "## Risks"]
    lines.extend(f"- [{risk.severity.upper()}] {risk.statement}" for risk in result.risks)
    lines += ["", "## Initial Solution Direction"]
    lines.extend(f"- **{item.service}:** {item.rationale}" for item in result.solution_direction)
    lines += ["", "## Next Steps"]
    lines.extend(f"- {step}" for step in result.decision.next_steps)
    return "\n".join(lines)


def analysis_to_json(result: AnalysisResult) -> bytes:
    return json.dumps(result.model_dump(mode="json"), indent=2).encode("utf-8")


def analysis_to_docx(result: AnalysisResult) -> bytes:
    document = Document()
    document.add_heading(f"{result.profile.customer_name} — AI Opportunity Brief", 0)
    document.add_paragraph(f"Use case: {result.profile.use_case}")
    document.add_paragraph(f"Discovery coverage: {result.coverage.score}%")
    document.add_paragraph(f"Readiness score: {result.readiness_score}%")
    document.add_paragraph(f"Recommended action: {result.decision.recommendation}")

    document.add_heading("Discovery Gaps", level=1)
    for gap in result.gaps:
        document.add_paragraph(gap.question, style="List Bullet")

    document.add_heading("Risks", level=1)
    for risk in result.risks:
        document.add_paragraph(
            f"{risk.statement} Mitigation: {risk.mitigation}", style="List Bullet"
        )

    document.add_heading("Initial Solution Direction", level=1)
    for item in result.solution_direction:
        document.add_paragraph(f"{item.service}: {item.rationale}", style="List Bullet")

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
